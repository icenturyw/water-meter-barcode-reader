from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageChops, ImageEnhance, ImageFilter, ImageOps, UnidentifiedImageError
import zxingcpp


SUPPORTED_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp",
}

STATUS_AUTO = "自动通过"
STATUS_REVIEW = "待复核"
STATUS_UNREAD = "未识别"
STATUS_MANUAL = "人工确认"


@dataclass(slots=True)
class ScanSettings:
    exact_length: int | None = 14
    digits_only: bool = True
    prefix: str = ""

    def validation_error(self, value: str) -> str:
        if self.digits_only and not value.isdigit():
            return "不是纯数字"
        if self.exact_length and len(value) != self.exact_length:
            return f"位数为 {len(value)}，期望 {self.exact_length} 位"
        if self.prefix and not value.startswith(self.prefix):
            return f"不符合前缀 {self.prefix}"
        return ""


@dataclass(slots=True)
class BarcodeCandidate:
    text: str
    format: str
    position: tuple[tuple[int, int], ...] = ()


@dataclass(slots=True)
class ScanRecord:
    source_path: Path
    sequence: int = 0
    meter_no: str = ""
    status: str = STATUS_UNREAD
    reason: str = ""
    candidates: list[BarcodeCandidate] = field(default_factory=list)
    capture_time: str = ""
    duplicate: bool = False
    manually_edited: bool = False
    review_time: str = ""

    @property
    def file_name(self) -> str:
        return self.source_path.name

    @property
    def raw_values(self) -> str:
        return " | ".join(candidate.text for candidate in self.candidates)

    @property
    def barcode_formats(self) -> str:
        return " | ".join(dict.fromkeys(candidate.format for candidate in self.candidates))

    @property
    def method(self) -> str:
        return "人工录入" if self.manually_edited else ("条码解码" if self.candidates else "")

    @property
    def is_exception(self) -> bool:
        return self.status in {STATUS_REVIEW, STATUS_UNREAD} or self.duplicate


def discover_images(folder: Path, recursive: bool = True) -> list[Path]:
    iterator = folder.rglob("*") if recursive else folder.glob("*")
    return sorted(
        (path for path in iterator if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS),
        key=lambda path: str(path).casefold(),
    )


def _capture_time(image: Image.Image, path: Path) -> str:
    try:
        exif = image.getexif()
        value = exif.get(36867) or exif.get(306)
        if value:
            return str(value)
    except Exception:
        pass
    try:
        return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
    except OSError:
        return ""


def _position_tuple(position: object) -> tuple[tuple[int, int], ...]:
    try:
        return tuple(
            (int(point.x), int(point.y))
            for point in (
                position.top_left,
                position.top_right,
                position.bottom_right,
                position.bottom_left,
            )
        )
    except Exception:
        return ()


def _decode(image: Image.Image) -> list[BarcodeCandidate]:
    decoded = zxingcpp.read_barcodes(
        image,
        try_rotate=True,
        try_downscale=True,
        try_invert=True,
    )
    candidates: list[BarcodeCandidate] = []
    seen: set[str] = set()
    for barcode in decoded:
        text = barcode.text.strip()
        fmt = str(barcode.format)
        if not text or text in seen:
            continue
        seen.add(text)
        candidates.append(
            BarcodeCandidate(
                text=text,
                format=fmt,
                position=_position_tuple(barcode.position),
            )
        )
    return candidates


def _decode_fixed_code128(image: Image.Image) -> list[BarcodeCandidate]:
    """Decode checksum-valid Code 128 candidates from a difficult-image variant."""
    decoded = zxingcpp.read_barcodes(
        image,
        formats=zxingcpp.BarcodeFormat.Code128,
        try_rotate=True,
        try_downscale=True,
        try_invert=True,
        binarizer=zxingcpp.Binarizer.FixedThreshold,
        return_errors=False,
    )
    candidates: list[BarcodeCandidate] = []
    seen: set[str] = set()
    for barcode in decoded:
        text = barcode.text.strip()
        if not barcode.valid or not text or text in seen:
            continue
        seen.add(text)
        candidates.append(BarcodeCandidate(text=text, format=str(barcode.format)))
    return candidates


def _focus_regions(gray: Image.Image) -> list[Image.Image]:
    """Return a small, deterministic set of overlapping regions for hard photos."""
    width, height = gray.size
    crop_width = max(1, round(width * 0.85))
    crop_height = max(1, round(height * 0.38))
    x_starts = (0, max(0, width - crop_width))
    max_y = max(0, height - crop_height)
    ordered_y_indexes = (2, 3, 1, 4, 0)
    y_starts = tuple(round(max_y * index / 4) for index in ordered_y_indexes)

    regions: list[Image.Image] = []
    boxes: set[tuple[int, int, int, int]] = set()
    for y in y_starts:
        for x in x_starts:
            box = (x, y, min(width, x + crop_width), min(height, y + crop_height))
            if box in boxes:
                continue
            boxes.add(box)
            regions.append(gray.crop(box))
    return regions


def _decode_difficult_code128(image: Image.Image, settings: ScanSettings) -> list[BarcodeCandidate]:
    """Bounded fallback for small, rolled or mildly blurred phone photos.

    A value must satisfy the configured meter-number rules and be decoded from
    at least two transformed variants before it is returned for auto-acceptance.
    """
    gray = ImageOps.grayscale(image)
    if max(gray.size) > 3200:
        gray.thumbnail((3200, 3200), Image.Resampling.LANCZOS)

    evidence: dict[str, tuple[BarcodeCandidate, int]] = {}
    angles = (-16, -12, -20, -8, 0, 8, 12, 16, 20)

    def add_evidence(processed: Image.Image) -> list[BarcodeCandidate]:
        for candidate in _decode_fixed_code128(processed):
            if settings.validation_error(candidate.text):
                continue
            original, count = evidence.get(candidate.text, (candidate, 0))
            evidence[candidate.text] = (original, count + 1)
        return [candidate for candidate, count in evidence.values() if count >= 2]

    for region in _focus_regions(gray):
        blurred = region.filter(ImageFilter.GaussianBlur(radius=8))
        normalized = ImageOps.autocontrast(
            ImageChops.subtract(region, blurred, scale=1.0, offset=128)
        )
        for angle in angles:
            rotated = normalized.rotate(
                angle,
                resample=Image.Resampling.BICUBIC,
                expand=True,
                fillcolor=128,
            )
            confirmed = add_evidence(rotated)
            if confirmed:
                # Transformed coordinates do not map directly to the preview.
                return confirmed
    return []


def _decode_with_fallbacks(image: Image.Image, settings: ScanSettings) -> list[BarcodeCandidate]:
    candidates = _decode(image)
    if candidates:
        return candidates

    gray = ImageOps.grayscale(image)
    variants: list[Image.Image] = [
        ImageOps.autocontrast(gray, cutoff=1),
        ImageEnhance.Contrast(gray).enhance(1.8).filter(
            ImageFilter.UnsharpMask(radius=1.4, percent=170, threshold=3)
        ),
    ]
    if max(gray.size) < 2400:
        variants.append(
            ImageOps.autocontrast(gray.resize((gray.width * 2, gray.height * 2), Image.Resampling.LANCZOS))
        )

    for variant in variants:
        candidates = _decode(variant)
        if candidates:
            # A resized fallback no longer has coordinates matching the preview.
            if variant.size != image.size:
                for candidate in candidates:
                    candidate.position = ()
            return candidates

    # Phone photos sometimes have perspective/roll that a 90-degree rotation
    # search cannot correct. Only failed images pay for these extra attempts.
    deskew_source = ImageOps.autocontrast(gray, cutoff=1)
    deskew_source.thumbnail((2600, 2600), Image.Resampling.LANCZOS)
    for angle in (-15, -10, -5, 5, 10, 15):
        rotated = deskew_source.rotate(angle, resample=Image.Resampling.BICUBIC, expand=True, fillcolor=255)
        candidates = _decode(rotated)
        if candidates:
            for candidate in candidates:
                candidate.position = ()
            return candidates
    return _decode_difficult_code128(image, settings)


def scan_image(path: Path, settings: ScanSettings, sequence: int = 0) -> ScanRecord:
    path = Path(path)
    record = ScanRecord(source_path=path, sequence=sequence)
    try:
        with Image.open(path) as opened:
            record.capture_time = _capture_time(opened, path)
            image = ImageOps.exif_transpose(opened).convert("RGB")
        record.candidates = _decode_with_fallbacks(image, settings)
    except UnidentifiedImageError:
        record.reason = "文件不是可读取的图片或图片已损坏"
        return record
    except Exception as exc:
        record.reason = f"读取失败：{type(exc).__name__}"
        return record

    if not record.candidates:
        record.reason = "没有检测到可解码的条码；可能条码太小、倾斜或模糊，请重拍"
        return record

    valid = [candidate for candidate in record.candidates if not settings.validation_error(candidate.text)]
    if len(valid) == 1 and len(record.candidates) == 1:
        record.meter_no = valid[0].text
        record.status = STATUS_AUTO
        return record

    record.status = STATUS_REVIEW
    if len(record.candidates) > 1:
        record.reason = f"检测到 {len(record.candidates)} 个不同条码，请人工选择"
    else:
        record.reason = settings.validation_error(record.candidates[0].text) or "条码需人工复核"
    return record


def apply_manual_value(record: ScanRecord, value: str, settings: ScanSettings) -> str:
    value = value.strip()
    validation_error = settings.validation_error(value)
    record.meter_no = value
    record.status = STATUS_MANUAL
    record.manually_edited = True
    record.review_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    record.reason = f"人工确认；{validation_error}" if validation_error else "人工确认"
    return validation_error


def mark_duplicates(records: Sequence[ScanRecord]) -> None:
    grouped: dict[str, list[ScanRecord]] = {}
    for record in records:
        record.duplicate = False
        if record.meter_no:
            grouped.setdefault(record.meter_no, []).append(record)

    for duplicate_group in grouped.values():
        if len(duplicate_group) > 1:
            for record in duplicate_group:
                record.duplicate = True


def _excel_safe(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


EXPORT_HEADERS = [
    "批次/箱号",
    "序号",
    "表号",
    "状态",
    "原图文件名",
    "原图完整路径",
    "拍摄时间",
    "条码码制",
    "识别方式",
    "原始识别值",
    "异常原因",
    "是否重复",
    "是否人工修改",
    "复核时间",
]


def export_rows(records: Iterable[ScanRecord], batch: str) -> list[list[object]]:
    rows: list[list[object]] = []
    for record in records:
        values: list[object] = [
            batch,
            record.sequence,
            record.meter_no,
            record.status,
            record.file_name,
            str(record.source_path),
            record.capture_time,
            record.barcode_formats,
            record.method,
            record.raw_values,
            record.reason,
            "是" if record.duplicate else "否",
            "是" if record.manually_edited else "否",
            record.review_time,
        ]
        rows.append([_excel_safe(value) for value in values])
    return rows


def export_csv(records: Sequence[ScanRecord], target: Path, batch: str) -> None:
    target = Path(target)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(EXPORT_HEADERS)
        writer.writerows(export_rows(records, batch))


def export_xlsx(records: Sequence[ScanRecord], target: Path, batch: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    target = Path(target)
    workbook = Workbook()
    all_sheet = workbook.active
    all_sheet.title = "全部结果"
    exception_sheet = workbook.create_sheet("待复核及异常")

    all_rows = export_rows(records, batch)
    exception_records = [record for record in records if record.is_exception]
    exception_rows = export_rows(exception_records, batch)

    for sheet, rows in ((all_sheet, all_rows), (exception_sheet, exception_rows)):
        sheet.append(EXPORT_HEADERS)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 26
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B87")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row_number in range(2, sheet.max_row + 1):
            sheet.cell(row_number, 3).number_format = "@"
            sheet.cell(row_number, 3).data_type = "s"
            for column in (1, 3, 4, 12, 13):
                sheet.cell(row_number, column).alignment = Alignment(horizontal="center")
        widths = [16, 8, 22, 12, 26, 56, 21, 16, 14, 32, 42, 12, 14, 21]
        for column_index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + column_index)].width = width

    workbook.save(target)


def parse_exact_length(value: str) -> int | None:
    value = value.strip()
    if not value:
        return None
    if not re.fullmatch(r"\d+", value):
        raise ValueError("期望位数必须是正整数，或留空表示不限制")
    parsed = int(value)
    if not 1 <= parsed <= 128:
        raise ValueError("期望位数应在 1 到 128 之间")
    return parsed
