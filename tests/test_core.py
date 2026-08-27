from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageFilter
import zxingcpp

from core import (
    STATUS_AUTO,
    STATUS_REVIEW,
    ScanRecord,
    ScanSettings,
    apply_manual_value,
    _decode_difficult_code128,
    discover_images,
    export_xlsx,
    mark_duplicates,
    scan_image,
)


def make_code128(path: Path, value: str, angle: float = 0) -> None:
    barcode = zxingcpp.create_barcode(value, zxingcpp.BarcodeFormat.Code128)
    generated = zxingcpp.write_barcode_to_image(barcode, scale=4, add_hrt=True)
    image = Image.fromarray(generated).convert("RGB")
    if angle:
        image = image.rotate(angle, resample=Image.Resampling.BICUBIC, expand=True, fillcolor="white")
    canvas = Image.new("RGB", (max(1000, image.width + 200), max(700, image.height + 200)), "white")
    canvas.paste(image, ((canvas.width - image.width) // 2, (canvas.height - image.height) // 2))
    canvas.save(path, quality=88)


class CoreTests(unittest.TestCase):
    def test_scans_code128_and_preserves_leading_zero(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            image_path = Path(temp) / "水表 001.jpg"
            make_code128(image_path, "00123456789012")
            record = scan_image(image_path, ScanSettings(exact_length=14), sequence=1)
            self.assertEqual(record.status, STATUS_AUTO)
            self.assertEqual(record.meter_no, "00123456789012")
            self.assertEqual(record.barcode_formats, "Code 128")

    def test_rule_mismatch_requires_review(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            image_path = Path(temp) / "meter.jpg"
            make_code128(image_path, "12345678901234")
            record = scan_image(image_path, ScanSettings(exact_length=13), sequence=1)
            self.assertEqual(record.status, STATUS_REVIEW)
            self.assertEqual(record.meter_no, "")
            self.assertIn("期望 13 位", record.reason)

    def test_rotated_photo_is_decoded(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            image_path = Path(temp) / "倾斜 10 度.jpg"
            make_code128(image_path, "12345678901234", angle=10)
            record = scan_image(image_path, ScanSettings(exact_length=14), sequence=1)
            self.assertEqual(record.meter_no, "12345678901234")

    def test_difficult_code128_fallback_requires_repeatable_decode(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            image_path = Path(temp) / "difficult.jpg"
            make_code128(image_path, "87654321098765", angle=16)
            with Image.open(image_path) as opened:
                image = opened.convert("RGB").filter(ImageFilter.GaussianBlur(radius=0.7))
            candidates = _decode_difficult_code128(image, ScanSettings(exact_length=14))
            self.assertEqual([candidate.text for candidate in candidates], ["87654321098765"])

    def test_discovery_supports_chinese_paths_and_recursion(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / "第一批 水表"
            nested = folder / "一号箱"
            nested.mkdir(parents=True)
            make_code128(nested / "照片.JPG", "12345678901234")
            self.assertEqual(len(discover_images(folder, recursive=True)), 1)
            self.assertEqual(len(discover_images(folder, recursive=False)), 0)

    def test_duplicates_and_xlsx_text_format(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            settings = ScanSettings(exact_length=14)
            first = ScanRecord(source_path=folder / "a.jpg", sequence=1)
            second = ScanRecord(source_path=folder / "b.jpg", sequence=2)
            apply_manual_value(first, "00123456789012", settings)
            apply_manual_value(second, "00123456789012", settings)
            mark_duplicates([first, second])
            self.assertTrue(first.duplicate)
            self.assertTrue(second.duplicate)

            target = folder / "结果.xlsx"
            export_xlsx([first, second], target, "一号箱")
            workbook = load_workbook(target, data_only=False)
            sheet = workbook["全部结果"]
            self.assertEqual(sheet["C2"].value, "00123456789012")
            self.assertEqual(sheet["C2"].number_format, "@")
            self.assertEqual(workbook["待复核及异常"].max_row, 3)


if __name__ == "__main__":
    unittest.main()
